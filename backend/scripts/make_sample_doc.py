"""Generate the sample policy fixtures used to verify Stages 2-3.

The content deliberately mirrors PROJECT.md Section 11's example (Travel Policy,
Section 6.2, "six months of continuous service") so the retrieval checkpoint has
an unambiguous ground-truth answer to look for.

Usage:
    python -m scripts.make_sample_doc [--format docx|pdf|both] [--out DIR]
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK

from .make_sample_pdf import build_pdf
from .policy_content import APPENDIX, SECTIONS, TITLE


def _add_table(doc: Document, header: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for idx, name in enumerate(header):
        cell = table.rows[0].cells[idx]
        cell.text = ""
        cell.paragraphs[0].add_run(name).bold = True
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value


def build_docx(path: Path) -> Path:
    doc = Document()
    doc.add_heading(TITLE, level=0)

    for heading, paragraphs, tables in SECTIONS:
        doc.add_heading(heading, level=1)
        for text in paragraphs:
            doc.add_paragraph(text)
        for caption, header, rows in tables:
            doc.add_paragraph()
            para = doc.add_paragraph()
            para.add_run(caption).bold = True
            _add_table(doc, header, rows)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading(APPENDIX[0], level=1)
    for text in APPENDIX[1]:
        doc.add_paragraph(text)

    doc.save(path)
    return path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--format", choices=["docx", "pdf", "both"], default="both")
    parser.add_argument("--out", default="scripts/fixtures")
    args = parser.parse_args()

    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)

    if args.format in ("docx", "both"):
        print(build_docx(out_dir / "acme_travel_policy.docx"))
    if args.format in ("pdf", "both"):
        print(build_pdf(out_dir / "acme_travel_policy.pdf"))


if __name__ == "__main__":
    main()
